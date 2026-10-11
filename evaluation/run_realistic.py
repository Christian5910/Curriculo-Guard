"""Currículos legítimos gerados por ferramentas reais: nenhum deve gerar alerta.

Os PDFs sintéticos de `run_pdfs.py` vêm do mesmo PyMuPDF que o motor usa; aqui os PDFs
vêm de outros produtores, que escrevem o conteúdo de formas diferentes:
  reportlab    (precisa de `pip install reportlab`; roda na CI)
  LibreOffice  (se `soffice` estiver no PATH e python-docx instalado)
  wkhtmltopdf  (se estiver no PATH)
Ferramentas ausentes são ignoradas com um aviso.

Uso (na raiz do projeto):
    python -m evaluation.run_realistic            # tabela
    python -m evaluation.run_realistic --check    # sai com erro se algum PDF gerar alerta
"""
import shutil, subprocess, sys, tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import numpy as np
import pymupdf
from engine import analyze_pdf

JOBS = [('Analista de Dados Sênior', 'Itaú Unibanco', '2021 – atual',
         ['Construí dashboards em Power BI e Python para a diretoria, reduzindo o tempo de fechamento em 30%.',
          'Enviei relatórios semanais por e-mail para 12 gestores e mantive as regras de negócio e diretrizes de compliance documentadas.',
          'Liderei 4 analistas e conduzi entrevistas técnicas para novos candidatos.']),
        ('Engenheira de Dados', 'Magazine Luiza', '2018 – 2021',
         ['Desenvolvi pipelines em SQL e Airflow, processando 2 TB por dia.',
          'Projetei prompts e avaliações para um assistente interno de suporte com LLM, com foco em segurança e privacidade.']),
        ('Estagiária de TI', 'Prefeitura de Belo Horizonte', '2016 – 2018', ['Apoiei a infraestrutura e documentei instruções de trabalho.'])]
EDUCATION = [('Bacharelado em Sistemas de Informação', 'UFMG', '2012 – 2016'), ('Pós-graduação em Ciência de Dados', 'PUC Minas', '2019 – 2020')]
SKILLS = ['Python', 'SQL', 'Power BI', 'Airflow', 'Docker', 'Git', 'Estatística', 'Comunicação']


def reportlab_canvas(path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=A4); width, height = A4
    c.setFillColor(colors.HexColor('#1f4e79')); c.rect(0, height - 90, width, 90, fill=1, stroke=0)
    c.setFillColor(colors.white); c.setFont('Helvetica-Bold', 24); c.drawString(40, height - 55, 'Maria Silva')
    c.setFont('Helvetica', 11); c.drawString(40, height - 75, 'Analista de Dados Sênior · maria.silva@email.com')
    y = height - 130
    for section in ('Experiência', 'Formação', 'Habilidades'):
        c.setFillColor(colors.black); c.setFont('Helvetica-Bold', 14); c.drawString(40, y, section)
        c.setStrokeColor(colors.black); c.setLineWidth(1.5); c.line(40, y - 4, width - 40, y - 4); y -= 26
        for i in range(4):
            c.setFont('Helvetica', 10); c.setFillColor(colors.HexColor('#222222')); c.drawString(50, y, f'Item {i + 1}: atuação em projetos de dados e BI')
            c.setFillColor(colors.HexColor('#999999')); c.drawRightString(width - 40, y, '2019 – 2023'); y -= 16
        y -= 14
    c.setFont('Helvetica', 7); c.setFillColor(colors.HexColor('#888888')); c.drawString(40, 25, 'Currículo atualizado em 2026'); c.save()


def reportlab_platypus(path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    styles = getSampleStyleSheet(); story = [Paragraph('Maria Silva', styles['Title']), Paragraph('Experiência', styles['Heading2'])]
    for title, company, period, bullets in JOBS:
        story.append(Paragraph(f'<b>{title}</b> — {company} ({period})', styles['Normal']))
        story.extend(Paragraph('• ' + b, styles['BodyText']) for b in bullets); story.append(Spacer(1, 8))
    story.append(Paragraph('Formação', styles['Heading2']))
    table = Table([['Curso', 'Instituição', 'Período']] + [list(row) for row in EDUCATION], colWidths=[240, 120, 100])
    table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .5, colors.black), ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#dde6f0')),
                               ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f2f2f2')])]))
    story.extend([table, Spacer(1, 10), Paragraph('Habilidades: ' + ', '.join(SKILLS), styles['Normal'])])
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


def libreoffice(directory):
    from docx import Document
    from docx.shared import RGBColor
    doc = Document(); doc.add_heading('Maria Silva', 0); doc.add_paragraph('maria.silva@email.com · +55 31 99999-0000 · Belo Horizonte, MG')
    doc.sections[0].footer.paragraphs[0].text = 'Currículo atualizado em 2026'; doc.add_heading('Experiência', 1)
    for title, company, period, bullets in JOBS:
        doc.add_heading(f'{title} — {company} ({period})', 2)
        for bullet in bullets: doc.add_paragraph(bullet, style='List Bullet')
    doc.add_heading('Formação', 1); table = doc.add_table(rows=1, cols=3); table.style = 'Light Grid Accent 1'
    for cell, text in zip(table.rows[0].cells, ('Curso', 'Instituição', 'Período')): cell.text = text
    for row in EDUCATION:
        for cell, text in zip(table.add_row().cells, row): cell.text = text
    doc.add_heading('Habilidades', 1); doc.add_paragraph(', '.join(SKILLS))
    run = doc.add_paragraph().add_run('Datas em cinza claro: 2019 – 2023'); run.font.color.rgb = RGBColor(0xA6, 0xA6, 0xA6)
    source = Path(directory) / 'word.docx'; doc.save(source)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', str(directory), str(source)], check=True, capture_output=True, timeout=180)
    return Path(directory) / 'word.pdf'


def wkhtmltopdf(directory, name):
    photo = np.zeros((150, 120, 3), dtype=np.uint8)
    for y in range(150):
        for x in range(120): photo[y, x] = [180 - int(40 * y / 150), 140 - int(30 * x / 120), 120]
    pymupdf.Pixmap(pymupdf.csRGB, 120, 150, photo.tobytes(), False).save(str(Path(directory) / 'foto.png'))
    css = 'body{font-family:Arial;font-size:11pt;color:#222;margin:0}h2{border-bottom:2px solid #333;padding-bottom:2px;font-size:14pt}.date{color:#999;font-size:9pt}'
    jobs = ''.join(f"<h3>{t} — {e} <span class=date>{d}</span></h3><ul>" + ''.join(f'<li>{b}</li>' for b in bs) + '</ul>' for t, e, d, bs in JOBS)
    badges = ''.join(f'<span class=b>{s}</span>' for s in SKILLS)
    pages = {
        'simples': f"<html><head><meta charset=utf-8><style>{css}</style></head><body style='padding:30px'><h1>Maria Silva</h1><h2>Experiência</h2>{jobs}</body></html>",
        'moderno': f"<html><head><meta charset=utf-8><style>{css}.h{{background:linear-gradient(90deg,#14325a,#2e7bb4);color:#fff;padding:22px 30px}}.s{{float:left;width:200px;background:#14325a;color:#fff;padding:20px;min-height:900px}}.m{{margin-left:240px;padding:20px}}.b{{display:inline-block;background:#2e7bb4;color:#fff;border-radius:4px;padding:2px 8px;margin:2px;font-size:9pt}}img{{width:100px}}</style></head><body><div class=h><h1>Maria Silva</h1></div><div class=s><img src='foto.png'><h3>Habilidades</h3>{badges}</div><div class=m><h2>Experiência</h2>{jobs}</div></body></html>",
    }
    source = Path(directory) / f'{name}.html'; source.write_text(pages[name], encoding='utf-8')
    target = Path(directory) / f'wk_{name}.pdf'
    subprocess.run(['wkhtmltopdf', '--quiet', '--enable-local-file-access', str(source), str(target)], capture_output=True, timeout=180)
    return target


def generate(directory):
    """Gera os PDFs disponíveis; devolve [(rótulo, caminho)] e a lista de ferramentas ausentes."""
    produced, missing = [], []
    import importlib.util
    if importlib.util.find_spec('reportlab') is None: missing.append('reportlab')
    else:
        reportlab_canvas(Path(directory) / 'reportlab_canvas.pdf'); reportlab_platypus(Path(directory) / 'reportlab_platypus.pdf')
        produced += [('reportlab (canvas: faixa, réguas, rodapé)', Path(directory) / 'reportlab_canvas.pdf'),
                     ('reportlab (platypus: tabela com bordas)', Path(directory) / 'reportlab_platypus.pdf')]
    if shutil.which('soffice'):
        try: produced.append(('LibreOffice (docx com tabela e rodapé)', libreoffice(directory)))
        except Exception as exc: missing.append(f'LibreOffice ({type(exc).__name__})')
    else: missing.append('LibreOffice')
    if shutil.which('wkhtmltopdf'):
        for name in ('simples', 'moderno'): produced.append((f'wkhtmltopdf ({name})', wkhtmltopdf(directory, name)))
    else: missing.append('wkhtmltopdf')
    return produced, missing


def main(check=False):
    with tempfile.TemporaryDirectory() as tmp:
        produced, missing = generate(tmp); failures = 0
        print(f"{'documento':46s} {'caracteres':>10s}  resultado")
        for label, path in produced:
            result = analyze_pdf(path); good = not result['alerts']; failures += not good
            detail = '' if good else f"  alertas={sorted({a['type'] for a in result['alerts']})}"
            print(f"{label:46s} {len(result['text']):>10d}  {'ok' if good else 'FALHA'}{detail}")
    if missing: print('\nignorados (ferramenta ausente):', ', '.join(missing))
    return 1 if check and failures else 0


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'): sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    raise SystemExit(main(check='--check' in sys.argv))
